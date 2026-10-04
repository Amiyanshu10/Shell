import os
from pathlib import Path

import openpyxl


# ============================================================
# CONFIGURATION
# ============================================================

# Root of your repository
REPO_PATH = r"C:\path\to\Pyate22"

# Input Excel file
EXCEL_PATH = r"C:\path\to\input.xlsx"

# Sheet containing automation_script column
SHEET_NAME = "Sheet1"

# Column header containing script names
SCRIPT_COLUMN = "automation_script"

# Output column name
OUTPUT_COLUMN = "GitPath"


# ============================================================
# FUNCTIONS
# ============================================================

def build_repo_index(repo_path):
    """
    Scan repository once and create an index of all files
    and directories.

    This is much faster than scanning the entire repository
    separately for every Excel row.
    """

    index = {}

    for root, dirs, files in os.walk(repo_path):

        # Index directories
        for directory in dirs:
            name = directory.lower()
            index.setdefault(name, []).append(
                os.path.join(root, directory)
            )

        # Index files
        for file in files:
            name = file.lower()
            index.setdefault(name, []).append(
                os.path.join(root, file)
            )

    return index


def get_relative_path(full_path, repo_path):
    """
    Convert:
        C:\\abc\\xyz\\Pyate22\\tests\\test_xxx

    into:
        Pyate22/tests/test_xxx/
    """

    repo = Path(repo_path)
    full = Path(full_path)

    relative = full.relative_to(repo.parent)

    # Convert Windows \ to /
    path = relative.as_posix()

    # If result points to a .py file, return its parent folder.
    if full.is_file():
        path = relative.parent.as_posix()

    return path.rstrip("/") + "/"


def find_script(script_name, repo_index, repo_path):
    """
    Find a script/folder in repository.

    Supports:
        test_dst_1008.py
        test_dst_1008
    """

    if script_name is None:
        return "NOT FOUND"

    script_name = str(script_name).strip()

    if not script_name:
        return "NOT FOUND"

    candidates = [
        script_name.lower()
    ]

    # If Excel doesn't contain .py, also try .py
    if not script_name.lower().endswith(".py"):
        candidates.append((script_name + ".py").lower())

    matches = []

    for candidate in candidates:
        matches.extend(repo_index.get(candidate, []))

    # Remove duplicates
    matches = list(dict.fromkeys(matches))

    if not matches:
        return "NOT FOUND"

    if len(matches) == 1:
        return get_relative_path(matches[0], repo_path)

    # Multiple scripts/folders with same name
    paths = [
        get_relative_path(match, repo_path)
        for match in matches
    ]

    return "MULTIPLE MATCHES: " + " | ".join(paths)


def find_column(ws, column_name):
    """Find Excel column number using header name."""

    for cell in ws[1]:
        if cell.value and str(cell.value).strip() == column_name:
            return cell.column

    raise ValueError(
        f"Column '{column_name}' not found in Excel."
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print("Scanning repository...")

    repo_index = build_repo_index(REPO_PATH)

    print(
        f"Repository indexed. "
        f"{len(repo_index)} unique names found."
    )

    # Load Excel
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    script_col = find_column(ws, SCRIPT_COLUMN)

    # Add output column
    output_col = ws.max_column + 1
    ws.cell(row=1, column=output_col).value = OUTPUT_COLUMN

    found = 0
    not_found = 0
    multiple = 0

    # Start after header
    for row in range(2, ws.max_row + 1):

        script_name = ws.cell(
            row=row,
            column=script_col
        ).value

        if not script_name:
            continue

        result = find_script(
            script_name,
            repo_index,
            REPO_PATH
        )

        ws.cell(
            row=row,
            column=output_col
        ).value = result

        if result == "NOT FOUND":
            not_found += 1
            print(f"[NOT FOUND] {script_name}")

        elif result.startswith("MULTIPLE MATCHES"):
            multiple += 1
            print(f"[MULTIPLE]  {script_name}")

        else:
            found += 1
            print(f"[FOUND]     {script_name} -> {result}")

    # Save separate output so original Excel is safe
    input_path = Path(EXCEL_PATH)

    output_path = input_path.with_name(
        input_path.stem + "_with_paths.xlsx"
    )

    wb.save(output_path)

    print("\n======================================")
    print("Completed")
    print("======================================")
    print(f"Found            : {found}")
    print(f"Not found        : {not_found}")
    print(f"Multiple matches : {multiple}")
    print(f"Output           : {output_path}")


if __name__ == "__main__":
    main()
